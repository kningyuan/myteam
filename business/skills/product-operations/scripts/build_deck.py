#!/usr/bin/env python3
"""从 deck_brief.yaml 生成 .pptx（python-pptx 兜底路径）。"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

try:
    import yaml
except ImportError:
    yaml = None  # type: ignore

# 品牌色（pitch / report 共用基础样式）
COLOR_PITCH_TITLE = (0x1A, 0x56, 0xDB)
COLOR_REPORT_TITLE = (0x0F, 0x76, 0x6E)
COLOR_BODY = (0x33, 0x33, 0x33)
COLOR_MUTED = (0x66, 0x66, 0x66)


def _load_brief(path: Path) -> dict:
    text = path.read_text(encoding="utf-8")
    if yaml is not None:
        data = yaml.safe_load(text)
        if isinstance(data, dict):
            return data
    raise SystemExit("需要 PyYAML：pip install PyYAML")


def _rgb(r: int, g: int, b: int):
    from pptx.dml.color import RGBColor
    return RGBColor(r, g, b)


def build(brief: dict, out_path: Path) -> int:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        print("ERROR: 需要 python-pptx：pip install python-pptx", file=sys.stderr)
        return 1

    tpl = (brief.get("template") or "pitch").strip().lower()
    accent = COLOR_PITCH_TITLE if tpl == "pitch" else COLOR_REPORT_TITLE

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slides = brief.get("slides") or []
    if not slides:
        print("ERROR: brief 无 slides", file=sys.stderr)
        return 1

    for spec in slides:
        stype = (spec.get("type") or "content").strip().lower()
        if stype == "title":
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            if slide.shapes.title:
                slide.shapes.title.text = spec.get("title") or brief.get("title") or ""
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = spec.get("subtitle") or brief.get("subtitle") or ""
            if slide.shapes.title and slide.shapes.title.text_frame:
                for p in slide.shapes.title.text_frame.paragraphs:
                    p.font.size = Pt(40)
                    p.font.bold = True
                    p.font.color.rgb = _rgb(*accent)
        elif stype == "section":
            slide = prs.slides.add_slide(prs.slide_layouts[2])
            if slide.shapes.title:
                slide.shapes.title.text = spec.get("title") or ""
                for p in slide.shapes.title.text_frame.paragraphs:
                    p.font.size = Pt(36)
                    p.font.bold = True
                    p.font.color.rgb = _rgb(*accent)
        elif stype == "closing":
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            if slide.shapes.title:
                slide.shapes.title.text = spec.get("title") or "谢谢"
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = spec.get("subtitle") or ""
        else:  # content
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            if slide.shapes.title:
                slide.shapes.title.text = spec.get("title") or ""
                for p in slide.shapes.title.text_frame.paragraphs:
                    p.font.size = Pt(32)
                    p.font.bold = True
                    p.font.color.rgb = _rgb(*accent)
            body = slide.placeholders[1] if len(slide.placeholders) > 1 else None
            if body:
                tf = body.text_frame
                tf.clear()
                bullets = spec.get("bullets") or []
                for i, line in enumerate(bullets):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.text = str(line)
                    p.level = 0
                    p.font.size = Pt(20)
                    p.font.color.rgb = _rgb(*COLOR_BODY)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    print(f"OK: wrote {out_path} ({len(prs.slides)} slides)")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description="deck_brief.yaml → deck.pptx")
    ap.add_argument("brief", type=Path, help="deck_brief.yaml 路径")
    ap.add_argument("output", type=Path, help="输出 .pptx 路径")
    args = ap.parse_args()
    if not args.brief.is_file():
        print(f"ERROR: brief 不存在: {args.brief}", file=sys.stderr)
        return 1
    brief = _load_brief(args.brief)
    return build(brief, args.output)


if __name__ == "__main__":
    sys.exit(main())
