#!/usr/bin/env python3
"""验收 deck.pptx + 可选交付物 Markdown。"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path


def verify_pptx(pptx: Path, *, min_slides: int = 4, min_bytes: int = 8000) -> list[str]:
    errs: list[str] = []
    if not pptx.is_file():
        return [f"deck 不存在: {pptx}"]
    size = pptx.stat().st_size
    if size < min_bytes:
        errs.append(f"deck 过小 ({size} bytes < {min_bytes})")
    try:
        from pptx import Presentation
    except ImportError:
        print("WARN: 未安装 python-pptx，跳过页数校验", file=sys.stderr)
        return errs
    prs = Presentation(str(pptx))
    n = len(prs.slides)
    if n < min_slides:
        errs.append(f"页数不足 ({n} < {min_slides})")
    empty_titles = 0
    for slide in prs.slides:
        title = ""
        if slide.shapes.title and slide.shapes.title.text:
            title = slide.shapes.title.text.strip()
        if not title and slide.slide_layout.name not in ("Blank",):
            empty_titles += 1
    if empty_titles > 2:
        errs.append(f"过多无标题页 ({empty_titles})")
    return errs


def verify_deliverable(md: Path, pptx: Path) -> list[str]:
    errs: list[str] = []
    if not md.is_file():
        return errs
    text = md.read_text(encoding="utf-8")
    for sec in ("演示目标", "文稿结构", "产出文件"):
        if sec not in text:
            errs.append(f"交付物缺少章节: {sec}")
    if "deck.pptx" not in text:
        errs.append("交付物未提及 deck.pptx")
    rel = pptx.name
    if md.parent == pptx.parent and rel not in text:
        errs.append("产出文件章节应写明 deck.pptx 路径")
    return errs


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--deliverable", type=Path, default=None)
    ap.add_argument("--min-slides", type=int, default=4)
    args = ap.parse_args()

    errs = verify_pptx(args.pptx, min_slides=args.min_slides)
    if args.deliverable:
        errs.extend(verify_deliverable(args.deliverable, args.pptx))

    cover = args.pptx.parent / "evidence" / "cover.png"
    if not cover.is_file():
        print(f"WARN: 无封面截图 {cover}（非阻塞）", file=sys.stderr)

    if errs:
        for e in errs:
            print(f"FAIL: {e}", file=sys.stderr)
        return 1
    print(f"verify_deck: PASS ({args.pptx})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
